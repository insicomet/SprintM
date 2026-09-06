"""Выгрузка «лестницы» снеговой нагрузки из подборщика ИНСИ.

Готовит src/data/snowLadder.json — таблицу, по которой ИНСИ переводит
снеговую нагрузку в пару «снеговой район + коэффициент k блока банка».

Источник — лист «снегветер»:
  AM6:AN23  надбавка к нагрузке по типу покрытия
  AB5:AH53  ступени: порог нагрузки → район и k, отдельно для γn=1 и γn=0,8
            (AE — несущая способность выбранной пары, район × k)

Строки 54–58 того же диапазона — лишний блок, оставшийся от прошлой
версии таблицы; он ломает возрастающий порядок и в выгрузку не идёт.
"""

import json
import os
import openpyxl
import warnings
from openpyxl.utils import column_index_from_string as ci

warnings.filterwarnings("ignore")

ENGINE = os.environ.get(
    "SPRINTM_ENGINE",
    "/root/.claude/uploads/726a18de-9a86-57ac-9d50-9c356d0f76b4/"
    "03f80a5e-____________________________________________________9__12__15__18__21__24_________15.xlsx",
)
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src", "data")

wb = openpyxl.load_workbook(ENGINE, data_only=True)
ws = wb["снегветер"]

NEEDS_ENGINEER = "уточнить при расчете у главного конструктора"

supplement = {}
for r in range(6, 24):
    name = ws.cell(row=r, column=ci("AM")).value
    value = ws.cell(row=r, column=ci("AN")).value
    if name is not None and value is not None:
        supplement[name] = value

steps = []
for r in range(5, 54):
    threshold = ws.cell(row=r, column=ci("AB")).value
    if not isinstance(threshold, (int, float)):
        continue
    district1 = ws.cell(row=r, column=ci("AC")).value
    k1 = ws.cell(row=r, column=ci("AD")).value
    capacity = ws.cell(row=r, column=ci("AE")).value
    district08 = ws.cell(row=r, column=ci("AG")).value
    k08 = ws.cell(row=r, column=ci("AH")).value

    steps.append(
        {
            "порог_кПа": threshold,
            "район_γn1": None if district1 == NEEDS_ENGINEER else district1,
            "k_γn1": None if k1 == NEEDS_ENGINEER else k1,
            "несущая_кПа": capacity if isinstance(capacity, (int, float)) else None,
            "район_γn08": None if district08 == NEEDS_ENGINEER else district08,
            "k_γn08": None if k08 == NEEDS_ENGINEER else k08,
        }
    )

assert steps == sorted(steps, key=lambda s: s["порог_кПа"]), "ступени должны идти по возрастанию"

path = os.path.join(OUT, "snowLadder.json")
with open(path, "w", encoding="utf-8") as f:
    json.dump({"надбавка_покрытия": supplement, "ступени": steps}, f, ensure_ascii=False, indent=2)
    f.write("\n")
print("wrote", path, f"({len(steps)} ступеней, {len(supplement)} покрытий)")
